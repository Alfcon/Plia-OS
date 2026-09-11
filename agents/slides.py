from __future__ import annotations

import json
import logging
import re
import time
import uuid
from pathlib import Path

import httpx
from pptx import Presentation

logger = logging.getLogger(__name__)

_TIMEOUT = 60


def _output_dir() -> Path:
    from core.config import get_config
    d = Path(get_config().memory_dir) / "generated"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", (text or "").lower()).strip("_")
    return (s or "slides")[:40]


def _fallback(topic: str) -> dict:
    return {"title": topic, "slides": [{"heading": topic, "bullets": []}]}


def _expand_outline(topic: str, num_slides: int) -> dict:
    """One Ollama call expanding a topic into a slide outline. Never raises."""
    from core.config import get_config
    cfg = get_config()
    try:
        prompt = (
            f"Create a slide deck outline about: {topic}. Use at most {num_slides} content "
            'slides. Reply with ONLY JSON: {"title": "...", "slides": [{"heading": "...", '
            '"bullets": ["...", "..."]}]}. No prose.'
        )
        r = httpx.post(
            f"{cfg.ollama_url}/api/chat",
            json={
                "model": cfg.ollama_model,
                "messages": [{"role": "user", "content": prompt}],
                "stream": False,
            },
            timeout=_TIMEOUT,
        )
        text = r.json()["message"]["content"]
        raw = json.loads(text[text.index("{") : text.rindex("}") + 1])
    except Exception:
        logger.exception("Slide outline expansion failed")
        return _fallback(topic)

    title = str(raw.get("title") or topic).strip() or topic
    slides: list[dict] = []
    raw_slides = raw.get("slides") if isinstance(raw, dict) else None
    if not isinstance(raw_slides, list):
        raw_slides = []
    for s in raw_slides:
        if not isinstance(s, dict):
            continue
        heading = str(s.get("heading") or "").strip()
        if not heading:
            continue
        raw_bullets = s.get("bullets")
        bullets = [str(b).strip() for b in raw_bullets if str(b).strip()] if isinstance(raw_bullets, list) else []
        slides.append({"heading": heading, "bullets": bullets})
    slides = slides[:num_slides]
    if not slides:
        return _fallback(topic)
    return {"title": title, "slides": slides}


def generate_slides(topic: str, num_slides: int = 5) -> str:
    """Generate a .pptx deck about a topic. Never raises."""
    try:
        outline = _expand_outline(topic, num_slides)
        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = outline["title"]
        for s in outline["slides"]:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = s["heading"]
            body = slide.placeholders[1].text_frame
            bullets = s["bullets"] or [""]
            body.text = bullets[0]
            for b in bullets[1:]:
                body.add_paragraph().text = b
        path = _output_dir() / f"slides_{_slug(topic)}_{int(time.time())}_{uuid.uuid4().hex[:6]}.pptx"
        prs.save(str(path))
        return f"Saved slides to `{path}`."
    except Exception as exc:
        logger.exception("Slide generation failed")
        return f"Slide generation failed: {exc}"
