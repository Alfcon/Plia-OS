import json
from unittest.mock import MagicMock, patch


def _resp(content):
    r = MagicMock()
    r.json.return_value = {"message": {"content": content}}
    return r


OUTLINE = {
    "title": "Intro to CRDTs",
    "slides": [
        {"heading": "What", "bullets": ["Conflict-free", "Replicated"]},
        {"heading": "Why", "bullets": ["Offline-first"]},
        {"heading": "How", "bullets": ["Merge functions", "Monotonic state"]},
    ],
}


def test_expand_outline_parses_valid():
    from agents.slides import _expand_outline
    with patch("httpx.post", return_value=_resp(json.dumps(OUTLINE))):
        o = _expand_outline("Intro to CRDTs", 5)
    assert o["title"] == "Intro to CRDTs"
    assert len(o["slides"]) == 3
    assert o["slides"][0]["heading"] == "What"


def test_expand_outline_caps_num_slides():
    from agents.slides import _expand_outline
    with patch("httpx.post", return_value=_resp(json.dumps(OUTLINE))):
        o = _expand_outline("Intro to CRDTs", 2)
    assert len(o["slides"]) == 2          # capped


def test_expand_outline_malformed_falls_back():
    from agents.slides import _expand_outline
    with patch("httpx.post", return_value=_resp("not json")):
        o = _expand_outline("My Topic", 5)
    assert o["title"] == "My Topic"
    assert len(o["slides"]) == 1          # minimal fallback deck


def test_expand_outline_http_error_falls_back():
    from agents.slides import _expand_outline
    with patch("httpx.post", side_effect=RuntimeError("boom")):
        o = _expand_outline("My Topic", 5)
    assert o["title"] == "My Topic"
    assert o["slides"]                    # non-empty fallback


def test_generate_slides_writes_pptx(tmp_path):
    from agents import slides
    from pptx import Presentation
    with patch("agents.slides._expand_outline", return_value=OUTLINE), \
         patch("agents.slides._output_dir", return_value=tmp_path):
        out = slides.generate_slides("Intro to CRDTs", 5)
    assert "saved slides to" in out.lower()
    # find the written file and read it back
    files = list(tmp_path.glob("*.pptx"))
    assert len(files) == 1
    prs = Presentation(str(files[0]))
    assert len(prs.slides) == 1 + 3       # title slide + 3 content slides


def test_generate_slides_build_error_returns_message(tmp_path):
    from agents import slides
    with patch("agents.slides._expand_outline", return_value=OUTLINE), \
         patch("agents.slides._output_dir", return_value=tmp_path), \
         patch("agents.slides.Presentation", side_effect=RuntimeError("bad pptx")):
        out = slides.generate_slides("Intro to CRDTs", 5)
    assert "failed" in out.lower()


def test_generate_slides_tool_registered(reset_registry):
    import sys
    from core.registry import set_loading_module, list_tools
    if "modules.slide_tools" in sys.modules:
        del sys.modules["modules.slide_tools"]
    set_loading_module("slide_tools")
    try:
        import modules.slide_tools  # noqa: F401
    finally:
        set_loading_module("")
    assert "generate_slides" in list_tools()
